"""Update the submission deck with reproducible proof-of-concept results.

    python scripts/update_deck.py                       # -> ShelfTwin.pptx
    python scripts/update_deck.py --results results/kaggle

Always starts from deck/ShelfTwin_base.pptx (the original deck), so it is safe to re-run when
new Kaggle results arrive. Numbers come from files in this repository:
  sim/results.json                     python sim/simulate.py 20  (baselines vs engine)
  results/kaggle/engine_analysis.json  engine tuned on held-out simulator seeds + error analysis
  results/kaggle/eda/*.json|png        EDA of SKU-110K and RP2K (Kaggle session 1)
  results/kaggle/final_report.json,
  results/kaggle/exp_*.json, ...       trained-model results (Kaggle session 2), shown as
                                       "pending" until they exist
"""
from __future__ import annotations

import argparse
import copy
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_AXIS_CROSSES, XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
REPO_URL = "https://github.com/Purv-Kabaria/shelftwin"
KAGGLE_URL = "https://www.kaggle.com/code/purv2k2/notebook77852dc329"
INK = RGBColor(0x14, 0x14, 0x14)
GREY = RGBColor(0xBF, 0xBF, 0xBF)
MID = RGBColor(0x7F, 0x7F, 0x7F)
LINE = RGBColor(0x8C, 0x8C, 0x8C)
PENDING = "pending"


# ============================================================================ small helpers
def load(path: Path):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def pct(x, d=1):
    return PENDING if x is None else ("100%" if round(x * 100, d) == 100 else f"{x * 100:.{d}f}%")


def num(x, d=2):
    return PENDING if x is None else f"{x:,.{d}f}"


def shape_by_name(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def set_runs(paragraph, text: str):
    """Replace a paragraph's text but keep the formatting of its first run."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run().text = text
        return paragraph.runs[0]
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return runs[0]


def set_paragraphs(text_frame, items: list[str]):
    """Set one paragraph per item, cloning the first paragraph's formatting (bullets etc.)."""
    paras = text_frame.paragraphs
    proto = copy.deepcopy(paras[0]._p)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    set_runs(text_frame.paragraphs[0], items[0])
    for it in items[1:]:
        el = copy.deepcopy(proto)
        text_frame._txBody.append(el)
        set_runs(text_frame.paragraphs[-1], it)


def add_link(run, url):
    run.hyperlink.address = url


def clone_into(shape, slide, x=None, y=None, w=None, h=None):
    """Copy a shape (keeps all formatting) onto another slide; optional new geometry (inches)."""
    el = copy.deepcopy(shape._element)
    for ext in el.iter(qn("p:extLst")):
        ext.getparent().remove(ext)
    slide.shapes._spTree.insert_element_before(el, "p:extLst")
    new = slide.shapes[-1]
    cnv = el.find(".//" + qn("p:cNvPr"))
    if cnv is not None:
        cnv.set("id", str(slide.shapes._next_shape_id))
    if x is not None:
        new.left = Inches(x)
    if y is not None:
        new.top = Inches(y)
    if w is not None:
        new.width = Inches(w)
    if h is not None:
        new.height = Inches(h)
    return new


def move_slide(prs, slide, index: int):
    lst = prs.slides._sldIdLst
    el = next(e for e in list(lst) if prs.part.related_part(e.get(qn("r:id"))) is slide.part)
    lst.remove(el)
    lst.insert(index, el)


# ============================================================================ tables (cloned style)
def build_table(template_frame, slide, x, y, col_w: list[float], rows: list[list[str]], row_h: list[float],
                bold_first_col=True, bold_rows=(), font_pt=None):
    """A new table with exactly the look of `template_frame` (header fill, borders, fonts)."""
    frame = clone_into(template_frame, slide, x, y)
    tbl = frame._element.find(".//" + qn("a:tbl"))
    grid = tbl.find(qn("a:tblGrid"))
    trs = tbl.findall(qn("a:tr"))
    head_tr, body_tr = copy.deepcopy(trs[0]), copy.deepcopy(trs[1])
    for g in list(grid):
        grid.remove(g)
    for w in col_w:
        grid.append(grid.makeelement(qn("a:gridCol"), {"w": str(int(Inches(w)))}))
    for tr in trs:
        tbl.remove(tr)
    for r, values in enumerate(rows):
        tr = copy.deepcopy(head_tr if r == 0 else body_tr)
        tr.set("h", str(int(Inches(row_h[min(r, len(row_h) - 1)]))))
        tcs = tr.findall(qn("a:tc"))
        while len(tcs) < len(col_w):
            tr.append(copy.deepcopy(tcs[-1]))
            tcs = tr.findall(qn("a:tc"))
        for extra in tcs[len(col_w):]:
            tr.remove(extra)
        tbl.append(tr)
    frame.width = Inches(sum(col_w))
    table = frame.table
    for r, values in enumerate(rows):
        for c, v in enumerate(values):
            cell = table.cell(r, c)
            tf = cell.text_frame
            lines = str(v).split("\n")
            for p in tf.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            set_runs(tf.paragraphs[0], lines[0])
            for extra in lines[1:]:
                el = copy.deepcopy(tf.paragraphs[0]._p)
                tf._txBody.append(el)
                set_runs(tf.paragraphs[-1], extra)
            for p in tf.paragraphs:
                for rn in p.runs:
                    rn.font.bold = (r == 0) or (c == 0 and bold_first_col) or (r in bold_rows)
                    if font_pt:
                        rn.font.size = Pt(font_pt)
    return frame


# ============================================================================ native charts
def style_chart(chart, title: str, font_pt=9.5):
    chart.font.name = "Calibri"
    chart.font.size = Pt(font_pt)
    chart.font.color.rgb = INK
    chart.has_title = True
    tf = chart.chart_title.text_frame
    tf.text = title
    r = tf.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.name = Pt(10.5), True, "Calibri"
    r.font.color.rgb = INK


def bar_chart(slide, x, y, w, h, title, cats, values, colors, number_format='0.00', log=False,
              horizontal=True, axis_title=None):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("value", values)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    chart = gf.chart
    style_chart(chart, title)
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = number_format, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold, dl.font.name = Pt(9.5), True, "Calibri"
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
    va, ca = chart.value_axis, chart.category_axis
    va.has_major_gridlines = False
    va.visible = False
    ca.format.line.color.rgb = LINE
    ca.tick_labels.font.size = Pt(9.5)
    ca.has_major_gridlines = False
    if horizontal:
        ca.reverse_order = True  # first category on top
    if log:
        scaling = va._element.find(qn("c:scaling"))
        lb = scaling.makeelement(qn("c:logBase"), {"val": "10"})
        scaling.insert(0, lb)
        va.minimum_scale = 0.01          # values below 1 must not draw leftwards from 1
        va.crosses = XL_AXIS_CROSSES.MINIMUM
    if axis_title:
        va.has_title = True
        va.axis_title.text_frame.text = axis_title
    return gf


def line_chart(slide, x, y, w, h, title, cats, series: dict, colors, number_format="0.0"):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series.items():
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    chart = gf.chart
    style_chart(chart, title)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(9.5)
    plot = chart.plots[0]
    for s, c in zip(plot.series, colors):
        s.format.line.color.rgb = c
        s.format.line.width = Pt(2)
        s.smooth = False
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = c
        s.marker.format.line.color.rgb = c
    for i, (s, vals) in enumerate(zip(plot.series, series.values())):
        if i == 0:  # first series sits near zero: label only its last point (no overlaps with its line)
            dlab = s.points[len(vals) - 1].data_label
            dlab.position = XL_LABEL_POSITION.ABOVE
            dlab.text_frame.text = format(vals[-1], ".1f")
            dlab.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
            dlab.text_frame.paragraphs[0].runs[0].font.bold = True
        else:
            dl = s.data_labels
            dl.show_value = True
            dl.number_format, dl.number_format_is_linked = number_format, False
            dl.position = XL_LABEL_POSITION.ABOVE
            dl.font.size = Pt(9)
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xE6, 0xE6)
    va.tick_labels.font.size = Pt(9)
    va.format.line.fill.background()
    chart.category_axis.tick_labels.font.size = Pt(9.5)
    chart.category_axis.format.line.color.rgb = LINE
    return gf


# ============================================================================ text helpers
def caption(slide, template, x, y, w, h, text):
    tb = clone_into(template, slide, x, y, w, h)
    set_paragraphs(tb.text_frame, [text])
    return tb


def heading(slide, template, x, y, w, text):
    tb = clone_into(template, slide, x, y, w, 0.3)
    set_paragraphs(tb.text_frame, [text])
    return tb


def bullets(slide, template, x, y, w, h, items):
    tb = clone_into(template, slide, x, y, w, h)
    set_paragraphs(tb.text_frame, items)
    return tb


def kpi(slide, box_tpl, x, y, w, h, value, label, sub=None):
    box = clone_into(box_tpl, slide, x, y, w, h)
    box.text_frame.text = ""
    tb = slide.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.1), Inches(w - 0.3), Inches(0.55))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    r = tf.paragraphs[0].add_run()
    r.text = value
    r.font.size, r.font.bold, r.font.name = Pt(26 if value != PENDING else 18), True, "Calibri"
    r.font.color.rgb = INK if value != PENDING else MID
    tb2 = slide.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.68), Inches(w - 0.3), Inches(h - 0.75))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = tf2.margin_top = tf2.margin_bottom = 0
    r2 = tf2.paragraphs[0].add_run()
    r2.text = label
    r2.font.size, r2.font.name, r2.font.bold = Pt(10.5), "Calibri", True
    r2.font.color.rgb = INK
    if sub:
        p = tf2.add_paragraph()
        r3 = p.add_run()
        r3.text = sub
        r3.font.size, r3.font.name = Pt(9.5), "Calibri"
        r3.font.color.theme_color = MSO_THEME_COLOR.ACCENT_5


def picture(slide, path: Path, x, y, w=None, h=None):
    """Insert keeping the aspect ratio; fits inside (w, h) when both are given."""
    iw, ih = Image.open(path).size
    if w and h:
        if iw / ih > w / h:
            h2 = w * ih / iw
            y, h = y + (h - h2) / 2, h2
        else:
            w2 = h * iw / ih
            x, w = x + (w - w2) / 2, w2
    elif w:
        h = w * ih / iw
    else:
        w = h * iw / ih
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def new_slide(prs, title: str, after_index: int, title_tpl=None):
    layout = next(l for l in prs.slide_layouts if l.name == "CONTENT")
    s = prs.slides.add_slide(layout)
    ph = s.placeholders[100]
    if title_tpl is not None:  # same position and insets as the hand-placed titles on other slides
        ph.left, ph.top, ph.width, ph.height = title_tpl.left, title_tpl.top, title_tpl.width, title_tpl.height
        old = ph._element.txBody.find(qn("a:bodyPr"))
        old.getparent().replace(old, copy.deepcopy(title_tpl._element.txBody.find(qn("a:bodyPr"))))
    ph.text_frame.text = title
    r = ph.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.name = Pt(25), True, "Calibri"
    r.font.color.theme_color = MSO_THEME_COLOR.TEXT_1
    move_slide(prs, s, after_index + 1)
    return s


# ============================================================================ slide edits
def fmt_row(name, r):
    return [name, num(r["false_alerts_per_cam_hour"]), pct(r["precision"]), pct(r["recall"]),
            f"{r['median_delay_s']:.0f} s", f"{r['p95_delay_s']:.0f} s"]


def edit_slide12(prs, s12, sim: dict, engine: dict | None):
    """Prototype results: reproducible numbers from this repo's simulator, plus the tuned engine."""
    res = sim["results"]
    order = ["Single-frame", "Single-frame + person mask", "Majority vote 3 of 5 + mask", "ShelfTwin engine"]
    rows = [["Decision method", "False alerts / camera-hour", "Precision", "Recall", "Median delay", "95th pct delay"]]
    rows += [fmt_row(k if k != "ShelfTwin engine" else "ShelfTwin engine (spec defaults)", res[k]) for k in order]
    tuned = (engine or {}).get("tuned", {}).get("ShelfTwin engine")
    if tuned:
        rows.append(fmt_row("ShelfTwin engine (tuned)", tuned))
    old = next(s for s in s12.shapes if s.has_table)
    widths = [c.width / 914400 for c in old.table.columns]
    build_table(old, s12, old.left / 914400, old.top / 914400, widths, rows,
                row_h=[0.36] + [0.36] * (len(rows) - 1), bold_rows={len(rows) - 1})
    old._element.getparent().remove(old._element)
    events = res["ShelfTwin engine"]["true_events"]
    set_paragraphs(shape_by_name(s12, "Text 1").text_frame, [
        f"Same simulated observations for every method: {sim['n_runs']} runs x 8 hours, 24 slots per camera, "
        f"{events:,} true restock events. 'Tuned': parameters chosen on separate simulator seeds (100-109), "
        "reported on these seeds."])
    # chart: native, log scale
    pic = next(s for s in s12.shapes if s.name.startswith("Chart: false alerts"))
    x, y, w, h = (v / 914400 for v in (pic.left, pic.top, pic.width, pic.height))
    pic._element.getparent().remove(pic._element)
    cats = ["Single-frame", "Single-frame + mask", "Majority vote 3 of 5", "ShelfTwin (defaults)"]
    vals = [float(num(res[k]["false_alerts_per_cam_hour"]).replace(",", "")) for k in order]  # same rounding as the table
    cols = [GREY, GREY, GREY, MID]
    if tuned:
        cats.append("ShelfTwin (tuned)")
        vals.append(float(num(tuned["false_alerts_per_cam_hour"])))
        cols.append(INK)
    bar_chart(s12, x, y - 0.05, w, h + 0.4, "False alerts per camera-hour (log scale), 160 camera-hours",
              cats, vals, cols, number_format="0.00", log=True)
    set_paragraphs(shape_by_name(s12, "Text 7").text_frame, [
        "shelftwin/engine + sim/simulate.py: temporal engine, three baselines and the shelf simulator.",
        "Full pipeline in code: quality gate, YOLO11 pack+gap detector, geometry, DINOv2+ArcFace identity, "
        "planogram alignment, tasks, dashboard.",
        "56 automated tests; one Kaggle notebook trains, evaluates and error-analyses every model.",
    ])
    t10 = shape_by_name(s12, "Text 10").text_frame
    add_link(set_runs(t10.paragraphs[0], "github.com/Purv-Kabaria/shelftwin"), REPO_URL)
    t11 = shape_by_name(s12, "Text 11").text_frame
    r = set_runs(t11.paragraphs[0], "Run: pytest -q  ·  python sim/simulate.py 20. Training logs: ")
    r2 = t11.paragraphs[0].add_run()
    r2.text = "kaggle.com/code/purv2k2/notebook77852dc329"
    r2.font.size, r2.font.name = r.font.size, "Calibri"
    add_link(r2, KAGGLE_URL)
    t5 = shape_by_name(s12, "Text 5").text_frame
    runs = t5.paragraphs[0].runs
    runs[0].text = "Scope: "
    if len(runs) > 1:
        runs[1].text = ("these numbers test the decision logic on simulated detector outputs. "
                        "Trained models on real held-out data: next two slides.")
        for extra in runs[2:]:
            extra._r.getparent().remove(extra._r)


def edit_slide1(s1, engine: dict | None, sim: dict):
    tbl = next(s for s in s1.shapes if s.has_table).table
    tuned = (engine or {}).get("tuned", {}).get("ShelfTwin engine")
    base = sim["results"]["Single-frame"]["false_alerts_per_cam_hour"]
    fa = tuned["false_alerts_per_cam_hour"] if tuned else sim["results"]["ShelfTwin engine"]["false_alerts_per_cam_hour"]
    for r in range(len(tbl.rows)):
        if tbl.cell(r, 0).text_frame.text.strip() == "Prototype":
            set_runs(tbl.cell(r, 1).text_frame.paragraphs[0],
                     f"Working prototype; code and trained models are public. In simulation: {fa:.2f} false alerts "
                     f"per camera-hour vs {base:,.0f} for a single-frame detector, at full recall.")
    team = shape_by_name(s1, "Text 4").text_frame.paragraphs[0]
    r = set_runs(team, "Team ShelfTwin, SVNIT Surat   |   ")
    r2 = team.add_run()
    r2.text = "github.com/Purv-Kabaria/shelftwin"
    r2.font.size, r2.font.name, r2.font.bold = r.font.size, "Calibri", False
    add_link(r2, REPO_URL)


def edit_slide18(s18, have_models: bool):
    tbl = next(s for s in s18.shapes if s.has_table).table
    trained = "trained on SKU-110K (Kaggle)" if have_models else "training on SKU-110K (Kaggle)"
    emb = "trained on RP2K" if have_models else "training on RP2K"
    now = {
        "Temporal state engine": ("Built and tested; tuned on held-out simulator seeds", None),
        "Baselines and simulator": ("Built, with false-alert error analysis", None),
        "Pack and gap detector": (f"Built: YOLO11 {trained}; gap labels from shelf geometry",
                                  "Fine-tune on Indian-store footage; TensorRT on edge"),
        "Visibility mask": ("Built with a pretrained person segmenter", "Fine-tune for carts and baskets"),
        "Product identity": (f"Built: DINOv2 + ArcFace {emb}; FAISS; abstention",
                             "Indian pack catalogue (3-5 photos per SKU); OCR tie-break"),
        "Planogram matcher": ("Built and tested (row alignment, layout guard)", "Layout drafting from shelf tags; approval flow"),
        "Alerts and tasks": ("Built: priority, batching, escalation, camera verification; Telegram; en/hi/gu",
                             "Staff mobile app, WhatsApp channel"),
        "Twin dashboard": ("Built: live slot view, tasks, alert replay", "Store-wide, multi-camera view"),
        "Backroom routing": ("Rack location carried in every task", "Rack camera; pick location from stock"),
        "Time-to-empty": ("Built: Gamma-Poisson estimate with a range", "Use billing data when connected"),
    }
    for r in range(1, len(tbl.rows)):
        key = tbl.cell(r, 0).text_frame.text.strip()
        if key in now:
            a, b = now[key]
            run = set_runs(tbl.cell(r, 1).text_frame.paragraphs[0], a)
            run.font.bold = True
            if b:
                set_runs(tbl.cell(r, 2).text_frame.paragraphs[0], b)


# ============================================================================ new slides
def slide_real_data(prs, idx, tpl, results: Path, engine: dict | None):
    eda = load(results / "eda" / "eda_sku110k.json") or {}
    rp = load(results / "eda" / "eda_rp2k.json") or {}
    s = new_slide(prs, "Real data: what the data told us before training", idx, tpl["title"])
    small = eda.get("box_short_side_after_resize", {})
    gaps = eda.get("real_gaps_per_image", {})
    bpi = eda.get("boxes_per_image", {})
    att = (engine or {}).get("attribution", {})
    n_false = att.get("n_false")
    missed = att.get("cause_counts", {}).get("blocked_mask_missed")
    rows = [["Finding", "Measured", "Decision it drove"],
            ["Packs per shelf photo", f"mean {bpi.get('mean', 0):.0f}, max {bpi.get('max', 0):.0f}",
             "Keep up to 1000 detections; YOLO11 with NMS (YOLO26 stops at 300)"],
            ["Packs under 16 px at the training size",
             "\n".join(f"{k} px: {v['lt16'] * 100:.1f}%" for k, v in small.items()) or PENDING,
             "Drop 640 px; compare 896 and 1024 px"],
            ["Photos with a real shelf gap", f"{gaps.get('share_with_gap', 0) * 100:.0f}% ({gaps.get('mean', 0):.1f} per photo)",
             "Train and test the gap class on real gaps"],
            ["RP2K product photos", f"{rp.get('classes', 0):,} SKUs, {rp.get('images', 0) / 1000:.0f}k photos",
             "Open-set split by SKU; cap 40 photos per SKU"],
            ["Engine false alerts (simulation)",
             f"{missed} of {n_false} involve a missed person mask" if n_false else PENDING,
             "Improve the person/cart mask first"]]
    build_table(tpl["table"], s, 0.5, 1.05, [2.1, 1.75, 3.15], rows, row_h=[0.34, 0.48, 0.72, 0.48, 0.48, 0.48],
                font_pt=10.5)
    samples = results / "eda" / "eda_sku110k_samples.png"
    if samples.exists():
        import tempfile
        im = Image.open(samples)
        crop = Path(tempfile.mkdtemp()) / "samples.png"   # drop the figure title strip
        w0, h0 = im.size
        im.crop((0, int(h0 * 0.07), w0, h0)).save(crop)
        picture(s, crop, 0.5, 4.3, w=7.0, h=2.05)
    caption(s, tpl["caption"], 0.5, 6.4, 7.0, 0.5,
            "Real SKU-110K photos (Kaggle run). Green: annotated packs. Red: gaps found by our shelf-row rule; "
            "most are real empty spaces, some are structural (shelf ends).")
    if small:
        bar_chart(s, 7.75, 0.95, 5.08, 2.55, "Share of packs smaller than 16 px",
                  [f"{k} px" for k in small], [v["lt16"] for v in small.values()],
                  [GREY, MID, INK][:len(small)], number_format="0%", horizontal=False)
    sens = [r for r in (engine or {}).get("sensitivity", []) if r["param"] == "p_mask_miss"]
    if sens:
        line_chart(s, 7.75, 3.7, 5.08, 2.65, "False alerts per camera-hour vs person-mask miss rate",
                   [f"{r['value'] * 100:.0f}%" for r in sens],
                   {"ShelfTwin engine": [r["ShelfTwin engine"]["false_alerts_per_cam_hour"] for r in sens],
                    "Majority vote 3 of 5": [r["Majority vote 3 of 5 + mask"]["false_alerts_per_cam_hour"] for r in sens]},
                   [INK, GREY])
    caption(s, tpl["caption"], 7.75, 6.4, 5.08, 0.5,
            f"EDA on all {sum(eda.get('images', {}).values()):,} SKU-110K photos and "
            f"{rp.get('images', 0):,} RP2K photos; engine sweep on 4 x 4 h of simulation.")
    s.notes_slide.notes_text_frame.text = (
        "Sources: results/kaggle/eda/eda_sku110k.json, eda_rp2k.json, engine_analysis.json in the repo; "
        f"produced by the public Kaggle notebook {KAGGLE_URL}.")
    return s


def slide_models(prs, idx, tpl, results: Path):
    rep = load(results / "final_report.json") or {}
    e2e = load(results / "e2e_summary.json") or rep.get("e2e_demo") or {}
    emb = load(results / "embedder" / "embedder_report.json") or {}
    expd = load(results / "exp_detector.json") or {}
    expg = load(results / "exp_gap.json") or {}
    expe = load(results / "exp_embed.json") or {}
    s = new_slide(prs, "Trained models: results on held-out real data", idx, tpl["title"])
    pk = (rep.get("pack_detector") or {}).get("metrics", {})
    pg = rep.get("pack_gap_detector") or {}
    test = emb.get("test") or {}
    kpis = [
        (pct(pk.get("mAP50")) if pk else PENDING, "Pack detection mAP50",
         f"mAP50-95 {pct(pk.get('mAP50_95'))}; 2,920 unseen SKU-110K photos" if pk else "SKU-110K test split (2,920 photos)"),
        (pct((pg.get("gap_recall_at_p90") or {}).get("recall_at_p")) if pg else PENDING, "Gap recall at 90% precision",
         "real gaps on unseen shelf photos"),
        (pct(test.get("recall_at_1")) if test else PENDING, "Product identity Recall@1",
         f"SKUs never seen in training; accepted precision {pct(test.get('accepted_precision'))}" if test
         else "SKUs never seen in training (open set)"),
        ((f"{e2e.get('engine_false_alerts')} vs {e2e.get('single_frame_alerts')}" if "engine_false_alerts" in e2e else PENDING),
         "False alerts: ShelfTwin vs single-frame",
         f"end-to-end demo on a real photo; stockout caught in {e2e.get('detection_delay_s', 0):.0f} s"
         if e2e.get("detection_delay_s") is not None else "end-to-end demo on a real shelf photo"),
    ]
    for i, (v, lab, sub) in enumerate(kpis):
        kpi(s, tpl["box"], 0.5 + i * 3.14, 1.0, 2.94, 1.35, v, lab, sub)
    # experiments table
    rows = [["Experiment (subset, short schedule)", "Options compared", "Winner"]]
    if expd.get("runs"):
        best = expd["winner"]
        pretty = lambda n: n[2:].replace("y11", "YOLO11").replace("_nomosaic", " no-mosaic").replace("_", " ")
        opts = ", ".join(f"{pretty(r['name'])}: {r['mAP50_95']:.3f}" for r in expd["runs"])
        rows.append(["Detector: model x resolution x mosaic (mAP50-95)", opts,
                     f"{best['model'].replace('.pt', '')} @ {best['imgsz']} px"])
    else:
        rows.append(["Detector: model x resolution x mosaic", "YOLO11 n/s; 896 vs 1024 px; mosaic on/off", PENDING])
    if expg.get("runs"):
        opts = ", ".join(f"{r['name']} {r['gap_recall_at_p90'] or 0:.2f}" for r in expg["runs"])
        rows.append(["Gap labels (gap recall at P=0.9)", opts, expg["winner"]["name"].replace("_", " ")])
    else:
        rows.append(["Gap labels", "real gaps only vs + synthetic erasing vs shelf-back fill", PENDING])
    if expe.get("runs"):
        zs = next((z for z in expe.get("zero_shot", []) if "vit_small" in z.get("model", "")), {})
        opts = (f"zero-shot {zs.get('recall_at_1', 0):.3f}, " if zs else "") + \
            ", ".join(f"{r['name'].replace('_', ' ')} {r['recall_at_1']:.3f}" for r in expe["runs"])
        win = "zero-shot DINOv2" if expe.get("zero_shot_wins") else expe["winner"]["name"].replace("_", " ")
        rows.append(["Identity: fine-tuning depth (Recall@1)", opts, win])
    else:
        rows.append(["Identity: fine-tuning depth", "zero-shot vs head only vs last 2 / 4 blocks", PENDING])
    build_table(tpl["table"], s, 0.5, 2.65, [2.45, 3.1, 1.45], rows, row_h=[0.34, 0.62, 0.62, 0.62], font_pt=10)
    caption(s, tpl["caption"], 0.5, 5.0, 7.0, 0.45,
            "Each experiment ran on a small subset with a short schedule; winners were picked automatically "
            "(decision_*.json) and used for the final training on the full data.")
    errs = load(results / "final_embed_errors" / "identity_error_analysis.json") or \
        load(results / "exp_embed_errors" / "identity_error_analysis.json") or {}
    conf = errs.get("most_confused", [])[:2]
    items = ["Error analysis per model: misses by pack size and shelf density, counting error, "
             "confused SKU pairs, precision vs coverage."]
    if conf:
        items.append("Most confused pair: " + conf[0]["true"] + " vs " + conf[0]["predicted"] +
                     " (look-alike packs: the system abstains instead of guessing).")
    items.append("All numbers reproducible from the public Kaggle notebook and the repository.")
    bullets(s, tpl["bullets"], 0.5, 5.5, 7.0, 1.5, items)
    trace = results / "e2e_trace.png"
    if trace.exists():
        picture(s, trace, 7.75, 2.6, w=5.08, h=3.75)
        caption(s, tpl["caption"], 7.75, 6.45, 5.08, 0.45,
                "End-to-end run on a real shelf photo with every trained model: stockout, shoppers, glare, "
                "swapped products, refill.")
    else:
        heading(s, tpl["heading"], 7.75, 2.65, 5.08, "End-to-end demo")
        bullets(s, tpl["bullets"], 7.75, 3.0, 5.08, 2.0, [
            "Every frame runs the full pipeline: quality gate, person mask, trained pack+gap detector, "
            "identity, planogram alignment, temporal engine, tasks.",
            "Scripted episode on a real shelf photo: a product runs out, shoppers block the view, glare, "
            "two products are swapped, a refill is verified.",
            "Trace and video: results pending (training run in progress)."])
    s.notes_slide.notes_text_frame.text = (
        "Sources: results/kaggle/final_report.json, exp_detector.json, exp_gap.json, exp_embed.json, "
        f"embedder/embedder_report.json, e2e_summary.json. Training log: {KAGGLE_URL}")
    return s


# ============================================================================ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=str(ROOT / "deck" / "ShelfTwin_base.pptx"))
    ap.add_argument("--out", default=str(ROOT / "ShelfTwin.pptx"))
    ap.add_argument("--results", default=str(ROOT / "results" / "kaggle"))
    a = ap.parse_args()
    results = Path(a.results)
    sim = load(ROOT / "sim" / "results.json")
    engine = load(results / "engine_analysis.json")
    prs = Presentation(a.base)
    s1, s12, s18 = prs.slides[0], prs.slides[11], prs.slides[17]
    tpl = dict(table=next(s for s in s12.shapes if s.has_table), title=next(s for s in s12.shapes if s.is_placeholder),
               caption=shape_by_name(s12, "Text 1"), heading=shape_by_name(s12, "Text 2"),
               bullets=shape_by_name(s12, "Text 3"), box=shape_by_name(s12, "Repo box"))
    # new slides first (they clone from the unedited slide-12 templates)
    slide_real_data(prs, 11, tpl, results, engine)
    slide_models(prs, 12, tpl, results)
    edit_slide12(prs, s12, sim, engine)
    edit_slide1(s1, engine, sim)
    edit_slide18(s18, have_models=(results / "final_report.json").exists())
    prs.save(a.out)
    fix_content_types(Path(a.out))
    print("wrote", a.out, "with", len(prs.slides), "slides")


def fix_content_types(path: Path):
    """The base deck labels JPEGs 'image/jpg'; python-pptx re-declares them as per-part overrides.
    Normalise to the standard <Default Extension="jpg" ContentType="image/jpeg"/>."""
    import re
    import shutil
    import tempfile
    import zipfile
    tmp = Path(tempfile.mkdtemp()) / path.name
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                ct = data.decode("utf-8")
                ct = re.sub(r'<Override PartName="/[^"]+\.(?:jpg|jpeg)" ContentType="image/jpe?g"/>', "", ct)
                defaults = ""
                for ext in ("jpg", "jpeg"):
                    if f'Extension="{ext}"' not in ct:
                        defaults += f'<Default Extension="{ext}" ContentType="image/jpeg"/>'
                ct = ct.replace("<Default ", defaults + "<Default ", 1)
                data = ct.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(tmp), str(path))


if __name__ == "__main__":
    main()
